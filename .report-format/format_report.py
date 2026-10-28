# شرح الملف: يجهز شكل تقرير Word السابق ومراجعات صفحاته.
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pathlib import Path
root=Path(r'C:\Users\pcx\Desktop\Plugin_ERP\.report-format')
d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0); s.font.size=Pt(12 if name=='Normal' else 16 if name=='Title' else 13)
 s.element.get_or_add_rPr().append(OxmlElement('w:rtl'))
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.12

def p(text,style=None):
 x=d.add_paragraph(text,style); x.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 pp=x._p.get_or_add_pPr(); pp.append(OxmlElement('w:bidi'))
 for r in x.runs:
  r._r.get_or_add_rPr().append(OxmlElement('w:rtl')); r.font.name='Arial'; r._r.get_or_add_rPr().rFonts.set(qn('w:cs'),'Arial')
 return x
items=[
('توضيح تقييم المؤشرات','يلزم توضيح الأساس المستخدم في تقييم المؤشرات بعبارة «جيد»، وشرح المقصود بتصنيف «فرصة».',0),
('مراجعة فرق القيمة','يوجد فرق بقيمة 5 ريالات يحتاج إلى المراجعة وتحديد سبب الاختلاف.',1),
('إظهار الفروع في الرسم البياني','لا تظهر الفروع في التقرير بالشكل المتوقع، ويُرجى مراجعة عرضها في الرسم البياني.',2),
('توضيح مؤشر الكمية لكل فاتورة','سبب القيمة الظاهرة لمؤشر الكمية لكل فاتورة غير واضح، ويُرجى توضيح طريقة احتسابها.',3),
('توضيح فرصة المندوب','يلزم شرح المقصود بتصنيف «فرصة» في هذا المؤشر، مع إظهار اسم المندوب.',4),
('مراجعة عرض الفروع في الجدول','لا تظهر الفروع في الجدول بالشكل المتوقع، ويُرجى مراجعة البيانات المعروضة.',5),
('تحسين مسميات المؤشرات وطريقة الاحتساب','كلمة «الوقت» مكررة، ومعنى «تسرب» غير واضح. يُقترح احتساب المؤشر بناءً على المبيعات الاعتيادية وعدد الحركات، مع فصل العمليات الاستثنائية. ويمكن الإبقاء على الوضع الحالي مؤقتًا إذا تعذّر التعديل.',6),
('معالجة الصفحة الرابعة','توجد مشكلة في الصفحة الرابعة من التقرير، بينما تظهر البيانات في الصفحة الخامسة؛ ويُرجى مراجعة توزيع المحتوى وفواصل الصفحات.',7),
('التحقق من بيانات الآجل والتحصيل','قد يكون العرض مضللًا إذا لم تكن المبالغ مخصصة لفواتير الآجل. الأرقام غير مؤكدة وتحتاج إلى التحقق قبل الاعتماد عليها.',8),
('مراجعة بيانات المرتجعات','بيانات المرتجعات لا تظهر في التقرير، ويُرجى مراجعة استرجاعها وعرضها.',9)]
groups=[[0,1],[2,3,4],[5,6,7],[8,9]]
for page,group in enumerate(groups):
 if page: d.add_page_break()
 if page==0:
  p('ملاحظات مراجعة تقرير إجمالي المبيعات','Title')
  p('تتضمن المراجعة الملاحظات المطلوب توضيحها أو معالجتها في التقرير، مع إرفاق صورة لكل ملاحظة. وتحتاج أرقام الآجل وبيانات المرتجعات إلى المراجعة قبل الاعتماد عليها.')
 for idx in group:
  title,body,img=items[idx]
  x=p(f'{idx+1}  {title}','Heading 2'); x.paragraph_format.keep_with_next=True
  x=p(body); x.paragraph_format.keep_with_next=True
  im=Image.open(root/f'image{img}.png'); w,h=im.size
  maxh=3.0 if img in [8,9] else 2.55 if img==2 else 2.1
  width=min(6.75,w/h*maxh)
  x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(14)
  x.add_run().add_picture(str(root/f'image{img}.png'),width=Inches(width))
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); r._r.addnext(field)
d.core_properties.title='ملاحظات مراجعة تقرير إجمالي المبيعات'
d.save(r'C:\Users\pcx\Desktop\تقرير إجمالي المبيعات - منسق.docx')
